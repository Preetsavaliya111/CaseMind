import hashlib
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from sqlalchemy.orm import Session, joinedload, selectinload

from app.core.config import settings
from app.db.database import SessionLocal
from app.models.document import Document, DocumentChunk
from app.models.user import User
from app.services.content_access_service import apply_content_scope
from app.services.organization_unit_service import user_department_ids, user_team_ids
from app.services.audit_service import record_audit
from app.providers.embeddings import EmbeddingProviderError, get_embedding_provider
from app.providers.vector_store import VectorStoreError, delete_document_vectors, upsert_document_chunks


ALLOWED_MEDIA_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip", "application/octet-stream"},
}


def _query(db: Session, current_user: User):
    query = db.query(Document).options(joinedload(Document.uploaded_by), selectinload(Document.chunks))
    return apply_content_scope(query, db, current_user, Document, Document.uploaded_by_id)


def list_documents(db: Session, current_user: User, *, page: int, page_size: int, document_status: str | None) -> dict:
    query = _query(db, current_user)
    if document_status:
        query = query.filter(Document.status == document_status)
    total = query.count()
    items = query.order_by(Document.updated_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return {"items": items, "total": total, "page": page, "page_size": page_size, "pages": math.ceil(total / page_size) if total else 0}


def get_document(db: Session, current_user: User, document_id: UUID) -> Document:
    document = _query(db, current_user).filter(Document.id == document_id).first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return document


def _storage_target(organization_id: UUID, document_id: UUID, extension: str) -> Path:
    base = Path(settings.DOCUMENT_STORAGE_PATH).expanduser().resolve()
    target = (base / str(organization_id) / f"{document_id}{extension}").resolve()
    try:
        target.relative_to(base)
    except ValueError as exc:
        raise RuntimeError("Invalid document storage path") from exc
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _validate_content(extension: str, content: bytes) -> None:
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="File content is not a valid PDF")
    if extension == ".docx":
        try:
            from io import BytesIO

            with zipfile.ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ValueError
        except (zipfile.BadZipFile, ValueError) as exc:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="File content is not a valid DOCX document") from exc
    if extension in {".txt", ".md"}:
        if b"\x00" in content:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Text documents cannot contain binary data")
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Text documents must use UTF-8 encoding") from exc


async def create_document(db: Session, current_user: User, upload: UploadFile) -> Document:
    original_name = Path(upload.filename or "").name
    extension = Path(original_name).suffix.lower()
    if not original_name or extension not in ALLOWED_MEDIA_TYPES:
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Supported document types are PDF, TXT, Markdown, and DOCX")
    media_type = (upload.content_type or "application/octet-stream").lower()
    if media_type not in ALLOWED_MEDIA_TYPES[extension]:
        raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="The declared file type does not match the extension")

    content = bytearray()
    while chunk := await upload.read(1024 * 1024):
        content.extend(chunk)
        if len(content) > settings.MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail=f"Documents cannot exceed {settings.MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    if not content:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Document is empty")
    _validate_content(extension, bytes(content))

    team_ids = user_team_ids(db, current_user)
    department_ids = user_department_ids(db, current_user)
    document = Document(
        organization_id=current_user.organization_id,
        uploaded_by_id=current_user.id,
        team_id=team_ids[0] if team_ids else None,
        department_id=department_ids[0] if department_ids else None,
        visibility="team" if team_ids else "organization",
        original_filename=original_name[:255],
        extension=extension,
        media_type=media_type,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        storage_key="pending",
        status="uploaded",
    )
    db.add(document)
    db.flush()
    target = _storage_target(current_user.organization_id, document.id, extension)
    target.write_bytes(content)
    document.storage_key = str(target)
    record_audit(db, current_user, "document.uploaded", "document", document.id, {"filename": document.original_filename, "size_bytes": document.size_bytes})
    db.commit()
    return get_document(db, current_user, document.id)


def _extract_text(path: Path, extension: str) -> str:
    if extension in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig")
    if extension == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if extension == ".docx":
        return "\n\n".join(paragraph.text for paragraph in DocxDocument(str(path)).paragraphs if paragraph.text.strip())
    raise ValueError("Unsupported document type")


def _chunk_text(text: str, target_size: int = 1200, overlap: int = 160) -> list[tuple[str, int, int]]:
    normalized = re.sub(r"[ \t]+", " ", text.replace("\r\n", "\n")).strip()
    if not normalized:
        return []
    chunks: list[tuple[str, int, int]] = []
    start = 0
    while start < len(normalized):
        end = min(start + target_size, len(normalized))
        if end < len(normalized):
            boundary = max(normalized.rfind("\n", start + target_size // 2, end), normalized.rfind(". ", start + target_size // 2, end))
            if boundary > start:
                end = boundary + 1
        value = normalized[start:end].strip()
        if value:
            chunks.append((value, start, end))
        if end >= len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks


def process_document(document_id: UUID, db: Session | None = None) -> None:
    owns_session = db is None
    session = db or SessionLocal()
    document = session.query(Document).filter(Document.id == document_id, Document.archived_at.is_(None)).first()
    if document is None:
        if owns_session:
            session.close()
        return
    try:
        document.status = "processing"
        document.error_code = None
        document.error_message = None
        session.commit()

        text = _extract_text(Path(document.storage_key), document.extension)
        chunks = _chunk_text(text)
        if not chunks:
            raise ValueError("No readable text was found in the document")

        session.query(DocumentChunk).filter(DocumentChunk.document_id == document.id).delete()
        records = [
            DocumentChunk(
                organization_id=document.organization_id,
                document_id=document.id,
                chunk_index=index,
                content=content,
                char_start=start,
                char_end=end,
                token_estimate=max(1, math.ceil(len(content) / 4)),
            )
            for index, (content, start, end) in enumerate(chunks)
        ]
        session.add_all(records)
        document.chunk_count = len(records)
        document.extracted_characters = len(text)
        session.flush()
        session.commit()

        provider = get_embedding_provider()
        if provider is None:
            document.status = "ready_for_indexing"
            session.commit()
            return

        vectors: list[list[float]] = []
        for offset in range(0, len(records), 64):
            vectors.extend(provider.embed([record.content for record in records[offset:offset + 64]]))
        points = [
            {
                "id": str(record.id),
                "vector": vector,
                "payload": {
                    "organization_id": str(document.organization_id),
                    "source_type": "document",
                    "source_id": str(document.id),
                    "document_id": str(document.id),
                    "chunk_id": str(record.id),
                    "chunk_index": record.chunk_index,
                    "filename": document.original_filename,
                    "team_id": str(document.team_id) if document.team_id else None,
                    "department_id": str(document.department_id) if document.department_id else None,
                    "visibility": document.visibility,
                    "created_at": document.created_at.isoformat(),
                },
            }
            for record, vector in zip(records, vectors, strict=True)
        ]
        upsert_document_chunks(points)
        document.status = "indexed"
        document.indexed_at = datetime.now(timezone.utc)
        session.commit()
    except (EmbeddingProviderError, VectorStoreError) as exc:
        session.rollback()
        document = session.query(Document).filter(Document.id == document_id).first()
        if document:
            document.status = "failed"
            document.error_code = "indexing_unavailable"
            document.error_message = str(exc)[:500]
            session.commit()
    except Exception:
        session.rollback()
        document = session.query(Document).filter(Document.id == document_id).first()
        if document:
            document.status = "failed"
            document.error_code = "extraction_failed"
            document.error_message = "The document could not be processed. Verify that it is readable and try again."
            session.commit()
    finally:
        if owns_session:
            session.close()


def archive_document(db: Session, current_user: User, document_id: UUID) -> None:
    document = get_document(db, current_user, document_id)
    document.archived_at = datetime.now(timezone.utc)
    record_audit(db, current_user, "document.archived", "document", document.id, {"filename": document.original_filename})
    db.commit()
    try:
        Path(document.storage_key).unlink(missing_ok=True)
        delete_document_vectors(str(current_user.organization_id), str(document.id))
    except (OSError, VectorStoreError):
        pass
